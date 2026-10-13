"""
Конструктор программ ПК — параметрическая генерация .docx
Готовит программу на основе JSON-файла с данными.

Использование:
    python build_program.py program_data.json

JSON-структура — см. program_data_example.json в этой же папке
"""
import json
import sys
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn


def set_doc_style(doc):
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_after = Pt(4)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

    for level in range(1, 4):
        hs = doc.styles[f'Heading {level}']
        hs.font.name = 'Times New Roman'
        hs.font.color.rgb = RGBColor(0, 0, 0)
        hs.font.bold = True
        if level == 1:
            hs.font.size = Pt(16)
        elif level == 2:
            hs.font.size = Pt(14)
        else:
            hs.font.size = Pt(12)


def add_table(doc, headers, rows, header_bold=True):
    table_data = [headers] + rows
    cols = len(headers)
    t = doc.add_table(rows=len(table_data), cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for i, row_data in enumerate(table_data):
        for j, cell_text in enumerate(row_data):
            cell = t.cell(i, j)
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(str(cell_text))
            run.font.size = Pt(10)
            run.font.name = 'Times New Roman'
            if i == 0 and header_bold:
                run.bold = True
    return t


def build_program(data, output_path):
    doc = Document()
    set_doc_style(doc)

    # --- Титул ---
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        f'ДОПОЛНИТЕЛЬНАЯ ПРОФЕССИОНАЛЬНАЯ ПРОГРАММА\n'
        f'повышения квалификации\n'
        f'«{data['name']}»'
    )
    run.bold = True
    run.font.size = Pt(16)
    run.font.name = 'Times New Roman'

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(f'\n{data['city']}\n{data['year']}')
    r2.font.size = Pt(12)
    r2.font.name = 'Times New Roman'

    doc.add_page_break()

    # --- УЧЕБНЫЙ ПЛАН ---
    doc.add_heading('УЧЕБНЫЙ ПЛАН', level=1)
    doc.add_paragraph(f'Цель: {data['goal']}')
    doc.add_paragraph(f'Категория слушателей: {data['audience']}')

    doc.add_paragraph('Календарный учебный график:')
    for line in [
        f'Объем программы: {data['hours']} академических часов',
        f'Режим занятий: {data['schedule_format']}',
        f'Общая продолжительность: {data['duration']}',
    ]:
        doc.add_paragraph(line, style='List Bullet')

    # Таблица учебного плана
    lessons = data['lessons']
    table_rows = [
        [
            str(l['num']), l['title'], str(l['hours']),
            str(l['lectures']), str(l['practice']), l['control']
        ] for l in lessons
    ]
    total_lectures = sum(l['lectures'] for l in lessons)
    total_practice = sum(l['practice'] for l in lessons)
    table_rows.append(['', 'Итого', str(data['hours']),
                        str(total_lectures), str(total_practice), ''])

    add_table(doc,
              ['№', 'Наименование занятий', 'Всего часов',
               'Лекции', 'Практика', 'Форма контроля'],
              table_rows)

    doc.add_page_break()

    # --- I. Пояснительная записка ---
    doc.add_heading('I. Пояснительная записка', level=1)

    doc.add_heading('1.1. Цели программы', level=2)
    doc.add_paragraph(data['purpose'])

    doc.add_heading('1.2. Задачи программы', level=2)
    for task in data['tasks']:
        doc.add_paragraph(task, style='List Bullet')

    doc.add_heading('1.3. Актуальность программы', level=2)
    doc.add_paragraph(data['relevance'])

    doc.add_heading('1.4. Уникальность программы', level=2)
    for u in data['uniqueness']:
        doc.add_paragraph(u, style='List Bullet')

    # Профстандарт
    if 'professional_standard' in data:
        doc.add_heading('1.5. Связь с профессиональным стандартом', level=2)
        ps = data['professional_standard']
        doc.add_paragraph(ps.get('description', ''))
        if 'table' in ps:
            add_table(doc,
                      ['Трудовая функция', 'Трудовые действия', 'Формируется на занятиях'],
                      ps['table'])

    doc.add_heading('1.6. Практическая значимость', level=2)
    doc.add_paragraph(data['practical_value'])

    doc.add_heading('1.7. Структура программы', level=2)
    doc.add_paragraph(data['structure'])

    doc.add_heading('1.8. Организационно-педагогические условиях', level=2)
    for cond in data['conditions']:
        doc.add_paragraph(cond, style='List Bullet')

    doc.add_page_break()

    # --- II. Требования к уровню освоения ---
    doc.add_heading('II. Требования к уровню освоения содержания учебной программы', level=1)
    doc.add_paragraph('В результате освоения программы обучающийся должен:')

    doc.add_heading('Знать:', level=2)
    for item in data['must_know']:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_heading('Уметь и владеть:', level=2)
    for item in data['must_be_able']:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_page_break()

    # --- III. Календарный график ---
    doc.add_heading('III. Календарный график учебного процесса', level=1)
    graph_rows = [
        [f'Занятие {l['num']}', l['topics'], str(l['hours'])]
        for l in lessons
    ]
    add_table(doc, ['Занятие', 'Темы', 'Часов'], graph_rows)
    doc.add_page_break()

    # --- IV. Оценочные материалы ---
    doc.add_heading('IV. Оценочные материалы', level=1)
    doc.add_paragraph(data['current_control'])
    doc.add_paragraph(data['final_assessment'])
    for item in data['portfolio']:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_heading('Критерии оценивания итоговой практической работы', level=2)
    criteria_rows = [
        [str(i + 1), c['text'], c['grade']]
        for i, c in enumerate(data['criteria'])
    ]
    add_table(doc, ['№', 'Критерии', 'Оценка'], criteria_rows)
    doc.add_page_break()

    # --- V. Содержание программы ---
    doc.add_heading('V. Содержание программы', level=1)

    for lesson in lessons:
        doc.add_heading(
            f'Занятие {lesson['num']}. {lesson['title']} ({lesson['hours']} часа)',
            level=2
        )
        doc.add_paragraph('Учебно-тематический план')

        topics_rows = [
            [str(t['num']), t['title'], str(t['hours']),
             str(t['lectures']), str(t['practice'])]
            for t in lesson['topics']
        ]
        add_table(doc,
                  ['№', 'Тема', 'Всего часов', 'Лекции', 'Практика'],
                  topics_rows)

        doc.add_paragraph('Содержание тем:')
        for t in lesson['topics']:
            p = doc.add_paragraph()
            run = p.add_run(f'Тема {t["num"]}. {t["title"]}')
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            doc.add_paragraph(t['content'])

        if 'questions' in lesson and lesson['questions']:
            doc.add_heading('Контрольные вопросы:', level=3)
            for q in lesson['questions']:
                doc.add_paragraph(q, style='List Bullet')

        doc.add_paragraph('')

    doc.add_page_break()

    # --- VI. Литература ---
    doc.add_heading('VI. Учебно-методическое и информационное обеспечение программы', level=1)

    doc.add_heading('Нормативные документы', level=2)
    for s in data['regulatory_docs']:
        doc.add_paragraph(s, style='List Number')

    doc.add_heading('Интернет-ресурсы', level=2)
    for s in data['internet_resources']:
        doc.add_paragraph(s, style='List Number')

    # --- Сохранение ---
    doc.save(output_path)
    print(f'OK: {output_path}')


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print('Использование: python build_program.py <program_data.json>')
        print('Пример: python build_program.py "Выпускной проект/Нейросети в работе преподавателя/neural_networks.json"')
        sys.exit(1)

    json_path = sys.argv[1]
    if not os.path.exists(json_path):
        print(f'Файл не найден: {json_path}')
        sys.exit(1)

    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    output_dir = data.get('output_dir', '.')
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, f'{data["name"]}.docx')

    build_program(data, output_path)
