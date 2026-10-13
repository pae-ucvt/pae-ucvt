"""
Генерация PowerPoint-презентации проекта «Субагент-Методолог»
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# --- Color palette ---
BG = RGBColor(0xF5, 0xF0, 0xFF)        # light lavender bg
CARD = RGBColor(0xFF, 0xFF, 0xFF)       # white card
DARK = RGBColor(0x2D, 0x1B, 0x69)       # dark purple text
MID = RGBColor(0x5A, 0x4F, 0x8A)        # medium purple
TEXT_SEC = RGBColor(0x6B, 0x6B, 0x83)   # secondary text
BLUE = RGBColor(0x3A, 0x5F, 0xE8)       # blue
LILAC = RGBColor(0x9B, 0x72, 0xFF)      # lilac
GREEN = RGBColor(0x2E, 0x7D, 0x32)      # green
ORANGE = RGBColor(0xD4, 0x53, 0x00)     # orange
WHITE = RGBColor(0xFF, 0xFF, 0xFF)      # white (for text on dark)
LIGHT_CARD = RGBColor(0xEE, 0xE8, 0xFF) # light lilac card
STRONG_BG = RGBColor(0xE0, 0xD8, 0xFC)  # stronger lilac


def add_bg(slide, color=BG):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def shape(slide, left, top, w, h, fill_color, text="",
          f_size=12, f_color=DARK, bold=False, align=PP_ALIGN.LEFT):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill_color
    s.line.fill.background()
    s.shadow.inherit = False
    if text:
        tf = s.text_frame
        tf.word_wrap = True
        tf.auto_size = None
        tf.margin_left = Pt(14)
        tf.margin_right = Pt(14)
        tf.margin_top = Pt(8)
        tf.margin_bottom = Pt(8)
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = text
        run.font.size = Pt(f_size)
        run.font.color.rgb = f_color
        run.font.bold = bold
        run.font.name = "Segoe UI"
    return s


def box(slide, left, top, w, h, text, f_size=14, f_color=DARK,
        bold=False, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, w, h)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(f_size)
    run.font.color.rgb = f_color
    run.font.bold = bold
    run.font.name = "Segoe UI"


def circle_num(slide, left, top, num, bg=BLUE):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, Inches(0.55), Inches(0.55))
    s.fill.solid()
    s.fill.fore_color.rgb = bg
    s.line.fill.background()
    tf = s.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = str(num)
    run.font.size = Pt(16)
    run.font.color.rgb = WHITE
    run.font.bold = True
    run.font.name = "Segoe UI"


# ================================================================
# SLIDE 1 — TITLE
# ================================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(sl)

# decorative top bar
shape(sl, Inches(0), Inches(0), Inches(13.333), Inches(0.08), DARK)
shape(sl, Inches(0), Inches(0.08), Inches(4), Inches(0.04), LILAC)

box(sl, Inches(1), Inches(1.5), Inches(5), Inches(0.6),
    "ПРОЕКТ", f_size=16, f_color=BLUE, bold=True)
box(sl, Inches(1), Inches(2.2), Inches(11), Inches(1.5),
    "СУБАГЕНТ-МЕТОДОЛОГ\n@methodologist", f_size=44, f_color=DARK, bold=True)
box(sl, Inches(1), Inches(4.2), Inches(10), Inches(0.8),
    "AI-ассистент в Open Code для автоматической генерации\nпрограмм повышения квалификации, планов занятий и тестов",
    f_size=18, f_color=TEXT_SEC)
box(sl, Inches(1), Inches(6.2), Inches(5), Inches(0.5),
    "Портфолио-проект  |  2026", f_size=12, f_color=TEXT_SEC)

# accent line at bottom
shape(sl, Inches(1), Inches(6.8), Inches(3), Inches(0.04), LILAC)


# ================================================================
# SLIDE 2 — ПРОБЛЕМА И РЕШЕНИЕ
# ================================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(sl)
shape(sl, Inches(0), Inches(0), Inches(13.333), Inches(0.08), DARK)

box(sl, Inches(0.8), Inches(0.5), Inches(5), Inches(0.7),
    "Проблема и решение", f_size=28, f_color=DARK, bold=True)

# Problem card
shape(sl, Inches(0.8), Inches(1.5), Inches(5.5), Inches(2.8), CARD,
      "ПРОБЛЕМА\n\n"
      "Разработка ДПП ПК вручную занимает 3–5 дней:\n"
      "  • 7 разделов, таблицы, контрольные вопросы\n"
      "  • Требования ФГОС и профстандартов\n"
      "  • Каждый раз заново верстать Word-документ",
      f_size=12, f_color=TEXT_SEC)

# Solution card
shape(sl, Inches(7), Inches(1.5), Inches(5.5), Inches(2.8), LIGHT_CARD,
      "РЕШЕНИЕ\n\n"
      "Суб-агент @methodologist:\n"
      "  • Создаёт программу за 1 минуту\n"
      "  • Соблюдает жёсткие правила\n"
      "  • Генерирует .md + .docx автоматически",
      f_size=12, f_color=DARK)

# Number blocks with colored background + white text
nums = [
    ("36", "Часов\nпрограмма", BLUE),
    ("3", "Сценария\nработы", LILAC),
    ("1", "Минута\nна генерацию", BLUE),
    ("7", "Разделов\nв программе", LILAC),
]
for i, (num, desc, color) in enumerate(nums):
    x = Inches(0.8 + i * 3.1)
    s = shape(sl, x, Inches(4.8), Inches(2.8), Inches(2.2), color)
    tf = s.text_frame
    tf.clear()
    tf.margin_top = Pt(20)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = num
    run.font.size = Pt(42)
    run.font.color.rgb = WHITE
    run.font.bold = True
    run.font.name = "Segoe UI"
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    run2 = p2.add_run()
    run2.text = desc
    run2.font.size = Pt(12)
    run2.font.color.rgb = WHITE
    run2.font.name = "Segoe UI"


# ================================================================
# SLIDE 3 — АРХИТЕКТУРА
# ================================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(sl)
shape(sl, Inches(0), Inches(0), Inches(13.333), Inches(0.08), DARK)

box(sl, Inches(0.8), Inches(0.5), Inches(5), Inches(0.7),
    "Архитектура проекта", f_size=28, f_color=DARK, bold=True)

# Left — folder tree as visual blocks with clearer hierarchy
tree = [
    ("📁 methodologist/", STRONG_BG, DARK, True, Inches(0.5), Inches(1.4), Inches(5.5)),
    ("  📄 .opencode/agents/methodologist.md", LIGHT_CARD, DARK, False, Inches(1.0), Inches(2.0), Inches(5.0)),
    ("  📁 templates/", STRONG_BG, DARK, True, Inches(1.0), Inches(2.6), Inches(5.0)),
    ("    📄 program_template.md  +  program_sample.md", CARD, MID, False, Inches(1.5), Inches(3.2), Inches(5.0)),
    ("    📄 lesson_template.md  +  lesson_sample.md", CARD, MID, False, Inches(1.5), Inches(3.8), Inches(5.0)),
    ("    📄 test_template.md", CARD, MID, False, Inches(1.5), Inches(4.4), Inches(4.5)),
    ("  📄 convert_to_docx.py", LIGHT_CARD, BLUE, False, Inches(1.0), Inches(5.0), Inches(4.5)),
    ("  📄 README.md", LIGHT_CARD, TEXT_SEC, False, Inches(1.0), Inches(5.6), Inches(3.5)),
]
for text, bg_c, f_c, bld, x, y, w in tree:
    shape(sl, x, y, w, Inches(0.5), bg_c, text, f_size=10, f_color=f_c, bold=bld)

# Right — flow diagram
box(sl, Inches(7.5), Inches(1.4), Inches(5), Inches(0.5),
    "Как это работает", f_size=18, f_color=DARK, bold=True)

flow = [
    ("Пользователь вводит команду\n@methodologist создай программу...", BLUE),
    ("Агент читает образец\nиз templates/program_sample.md", BLUE),
    ("Копирует формат,\nподставляет новую тему", BLUE),
    ("Генерирует .md и запускает\nconvert_to_docx.py → .docx", BLUE),
    ("Готово: .md + .docx\nоткрывается в Word", GREEN),
]
for i, (desc, color) in enumerate(flow):
    y = Inches(2.1 + i * 1.0)
    circle_num(sl, Inches(7.5), y + Inches(0.1), i + 1, color)
    shape(sl, Inches(8.3), y, Inches(4.5), Inches(0.8), CARD, desc, f_size=11, f_color=DARK)


# ================================================================
# SLIDE 4 — СЦЕНАРИИ
# ================================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(sl)
shape(sl, Inches(0), Inches(0), Inches(13.333), Inches(0.08), DARK)

box(sl, Inches(0.8), Inches(0.5), Inches(5), Inches(0.7),
    "Сценарии работы", f_size=28, f_color=DARK, bold=True)

scenarios = [
    ("Сценарий 1", "Программа курса", BLUE,
     "Команда:\n@methodologist создай программу курса\nпо теме \"...\" на 36 часов, категория — ...\n\n"
     "Результат: ДПП ПК из 7 разделов\n"
     "Режим: 4 ч/день, 1-2 дня/нед, 4,5-9 нед\n"
     "Контрольные вопросы после каждой темы"),
    ("Сценарий 2", "План занятия", ORANGE,
     "Команда:\n@methodologist составь план занятия\nпо теме \"...\"\n\n"
     "Результат: хронометраж (таблица)\n"
     "6 этапов, активности, домашнее задание\n"
     "Длительность: 180 мин (4 ак. часа)"),
    ("Сценарий 3", "Тест", LILAC,
     "Команда:\n@methodologist сделай тест из 10\nвопросов по теме \"...\"\n\n"
     "Результат: 10 вопросов A-D\n"
     "ключ ответов, критерии оценки\n"
     "Смешанные типы: закрытые + открытые"),
]

for i, (title, sub, color, desc) in enumerate(scenarios):
    x = Inches(0.5 + i * 4.2)
    shape(sl, x, Inches(1.4), Inches(3.8), Inches(0.7), color, title,
          f_size=12, f_color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    shape(sl, x, Inches(2.1), Inches(3.8), Inches(0.6), STRONG_BG, sub,
          f_size=14, f_color=DARK, bold=True, align=PP_ALIGN.CENTER)
    shape(sl, x, Inches(2.8), Inches(3.8), Inches(4.0), CARD, desc, f_size=11)


# ================================================================
# SLIDE 5 — ПРАВИЛА
# ================================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(sl)
shape(sl, Inches(0), Inches(0), Inches(13.333), Inches(0.08), DARK)

box(sl, Inches(0.8), Inches(0.5), Inches(8), Inches(0.7),
    "Жёсткие правила (не нарушаются)", f_size=28, f_color=DARK, bold=True)

rules = [
    ("4 ч/день", "Режим: 4 часа в день,\n1–2 дня в неделю", BLUE),
    ("4,5–9 нед", "Продолжительность:\nот 4,5 до 9 недель", LILAC),
    ("Вопросы", "Контрольные вопросы\nпосле КАЖДОЙ темы (мин. 3)", GREEN),
    ("Источники", "Никаких выдуманных книг.\nТолько ФЗ, ФГОС, ссылки", ORANGE),
    ("30/70", "Лекции 30% / Практика 70%\nПриоритет практики", BLUE),
    ("Образец", "Копировать формат\nиз program_sample / lesson_sample", LILAC),
]

for i, (title, desc, color) in enumerate(rules):
    col = i % 3
    row = i // 3
    x = Inches(0.5 + col * 4.2)
    y = Inches(1.5 + row * 2.8)
    shape(sl, x, y, Inches(3.8), Inches(0.6), color, title,
          f_size=14, f_color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    shape(sl, x, y + Inches(0.7), Inches(3.8), Inches(1.7), CARD, desc, f_size=11)


# ================================================================
# SLIDE 6 — КОНВЕРТАЦИЯ
# ================================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(sl)
shape(sl, Inches(0), Inches(0), Inches(13.333), Inches(0.08), DARK)

box(sl, Inches(0.8), Inches(0.5), Inches(6), Inches(0.7),
    "Формат на выходе: .md → .docx", f_size=28, f_color=DARK, bold=True)

# Before
shape(sl, Inches(0.8), Inches(1.6), Inches(3.8), Inches(4.5), STRONG_BG,
      "", f_size=0)
box(sl, Inches(1.0), Inches(1.8), Inches(3.4), Inches(0.5),
    "Было (до проекта)", f_size=16, f_color=DARK, bold=True)
shape(sl, Inches(1.0), Inches(2.4), Inches(3.4), Inches(3.4), CARD,
      "Вручную писать 7 разделов\n"
      "Верстать таблицы в Word\n"
      "Форматировать вручную\n\n"
      "3–5 дней на программу\n"
      "Легко пропустить требования\n"
      "Ошибки в режиме / объёме",
      f_size=12, f_color=TEXT_SEC)

# Arrow
box(sl, Inches(5.1), Inches(3.2), Inches(1.5), Inches(0.8),
    "→", f_size=56, f_color=BLUE, bold=True, align=PP_ALIGN.CENTER)

# After
shape(sl, Inches(6.8), Inches(1.6), Inches(3.8), Inches(4.5), LIGHT_CARD,
      "", f_size=0)
box(sl, Inches(7.0), Inches(1.8), Inches(3.4), Inches(0.5),
    "Стало (с агентом)", f_size=16, f_color=GREEN, bold=True)
shape(sl, Inches(7.0), Inches(2.4), Inches(3.4), Inches(3.4), CARD,
      "Команда @methodologist\n"
      "Программа за 1 минуту\n"
      "Автоматический .md\n\n"
      "Автоконвертация в .docx\n"
      "Все правила соблюдены\n"
      "Контрольные вопросы есть",
      f_size=12, f_color=DARK)

# Converter explanation
shape(sl, Inches(11.2), Inches(1.6), Inches(2.0), Inches(4.5), STRONG_BG,
      "", f_size=0)
box(sl, Inches(11.3), Inches(1.8), Inches(1.8), Inches(0.5),
    "Конвертер", f_size=14, f_color=DARK, bold=True)
shape(sl, Inches(11.3), Inches(2.4), Inches(1.8), Inches(3.4), CARD,
      "convert_to_docx.py\n\n"
      "Запускается\nагентом\nавтоматически\n\n"
      "Форматирует:\nTimes New Roman\nтаблицы, списки",
      f_size=10, f_color=TEXT_SEC)


# ================================================================
# SLIDE 7 — ДЕМОНСТРАЦИЯ
# ================================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(sl)
shape(sl, Inches(0), Inches(0), Inches(13.333), Inches(0.08), DARK)

box(sl, Inches(0.8), Inches(0.5), Inches(5), Inches(0.7),
    "Демонстрация работы", f_size=28, f_color=DARK, bold=True)

steps = [
    ("1", "Открыть папку methodologist/ в новой сессии Open Code", BLUE),
    ("2", "@methodologist создай программу курса по теме \"Нейросети в работе преподавателя\" на 36 часов, категория — преподаватели СПО", LILAC),
    ("3", "Агент читает образец → генерирует .md → запускает конвертер", BLUE),
    ("4", "Подтвердить конвертацию (агент запросит разрешение)", LILAC),
    ("5", "Готово: .md + .docx  —  открывайте в Word", GREEN),
]

for i, (num, desc, color) in enumerate(steps):
    y = Inches(1.5 + i * 1.05)
    circle_num(sl, Inches(0.8), y + Inches(0.05), num, color)
    bg = CARD if i % 2 == 0 else LIGHT_CARD
    shape(sl, Inches(1.6), y, Inches(10.8), Inches(0.8), bg, desc, f_size=11)

# Tip
shape(sl, Inches(0.8), Inches(6.8), Inches(11.6), Inches(0.5), STRONG_BG,
      "Совет: для показа портфолио — откройте git log и папку methodologist/ в редакторе",
      f_size=11, f_color=DARK, align=PP_ALIGN.CENTER)


# ================================================================
# SLIDE 8 — ИТОГ
# ================================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(sl)
shape(sl, Inches(0), Inches(0), Inches(13.333), Inches(0.08), DARK)

box(sl, Inches(0.8), Inches(0.5), Inches(5), Inches(0.7),
    "Итоги проекта", f_size=28, f_color=DARK, bold=True)

summary = [
    ("Что создано", "Суб-агент @methodologist:\nконфиг + шаблоны + образцы + конвертер"),
    ("Сценариев", "3: программа курса, план занятия, тест"),
    ("Файлов в проекте", "10 файлов, ~1550 строк"),
    ("Формат на выходе", ".md (Markdown) + .docx (Word) — автоматически"),
    ("Режим", "4 ч/день, 1-2 дня/нед, 4,5-9 нед, 36 ч"),
    ("Контроль качества", "Вопросы после каждой темы, реальные источники"),
]

for i, (key, val) in enumerate(summary):
    y = Inches(1.4 + i * 0.9)
    color = BLUE if i % 2 == 0 else LILAC
    shape(sl, Inches(0.8), y, Inches(3.2), Inches(0.7), color, key,
          f_size=12, f_color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    shape(sl, Inches(4.2), y, Inches(8.3), Inches(0.7), CARD, val, f_size=11)

# Bottom line
shape(sl, Inches(0.8), Inches(6.9), Inches(6), Inches(0.04), LILAC)


# ================================================================
# SAVE
# ================================================================
output = (
    r"C:\Users\user\Downloads\for_Cloude Code\prepord-card\methodologist"
    r"\Субагент-Методолог_презентация.pptx"
)
prs.save(output)
print(f"OK -> {output}")
