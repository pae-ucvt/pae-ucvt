"""
Генерация обобщающего конспекта-инструкции по проекту «Субагент-Методолог»
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()

# --- styles ---
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)
style.paragraph_format.line_spacing = 1.15
style.paragraph_format.space_after = Pt(4)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

for level in range(1, 4):
    hs = doc.styles[f"Heading {level}"]
    hs.font.name = "Times New Roman"
    hs.font.color.rgb = RGBColor(0, 0, 0)
    hs.font.bold = True
    hs.font.size = Pt({1: 16, 2: 14, 3: 12}[level])

# ---- TITLE ----
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("ПРОЕКТ «СУБАГЕНТ-МЕТОДОЛОГ»\nОбобщающий конспект-инструкция")
run.bold = True
run.font.size = Pt(16)
run.font.name = "Times New Roman"

doc.add_paragraph("")

# ---- 1. НАЗНАЧЕНИЕ ----
doc.add_heading("1. Назначение проекта", level=1)
doc.add_paragraph(
    "Суб-агент @methodologist в Open Code автоматически создаёт документы "
    "для дополнительного профессионального образования — повышения квалификации (ДПП ПК):"
)
items = [
    "Программа курса (36 ч, 7 разделов)",
    "План одного занятия (хронометраж, активности, ДЗ)",
    "Тест (10 вопросов + ключ + критерии)",
]
for it in items:
    doc.add_paragraph(it, style="List Bullet")

# ---- 2. СТРУКТУРА ПРОЕКТА ----
doc.add_heading("2. Структура проекта", level=1)

struct_data = [
    ["Папка / Файл", "Назначение"],
    ["methodologist/", "Корневая папка проекта"],
    [".opencode/agents/methodologist.md", "Конфигурация и инструкция суб-агента"],
    ["templates/program_template.md", "Шаблон программы курса"],
    ["templates/program_sample.md", "Образец программы (для копирования агентом)"],
    ["templates/lesson_template.md", "Шаблон плана занятия"],
    ["templates/lesson_sample.md", "Образец плана занятия (для копирования агентом)"],
    ["templates/test_template.md", "Шаблон теста"],
    ["scripts/convert_to_docx.py", "Конвертер .md → .docx (запускается агентом автоматически)"],
    ["README.md", "Документация для портфолио"],
]

t = doc.add_table(rows=len(struct_data), cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.style = "Table Grid"
for i, row_data in enumerate(struct_data):
    for j, cell_text in enumerate(row_data):
        cell = t.cell(i, j)
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(cell_text)
        run.font.size = Pt(10)
        run.font.name = "Times New Roman"
        run.bold = (i == 0)

# ---- 3. РАЗГРАНИЧЕНИЕ ----
doc.add_heading("3. Что делаете вы и что делаю я", level=1)

roles_data = [
    ["Этап", "Вы (пользователь)", "Я (AI-ассистент)"],
    [
        "Подготовка материалов",
        "Предоставляете файл-референс ДПП ПК (.doc/.docx) для анализа структуры",
        "Извлекаю текст из .doc/.docx, анализирую структуру, создаю шаблоны",
    ],
    [
        "Настройка агента",
        "Открываете папку methodologist/ в новой сессии Open Code",
        "Создаю конфиг .opencode/agents/methodologist.md с инструкцией и правами",
    ],
    [
        "Генерация программы",
        "Пишете команду: @methodologist создай программу...",
        "Читаю образец program_sample.md, копирую формат, подставляю новую тему, сохраняю .md + .docx",
    ],
    [
        "Генерация плана занятия",
        "Пишете команду: @methodologist составь план занятия...",
        "Читаю образец lesson_sample.md, копирую формат, подставляю новую тему, сохраняю .md + .docx",
    ],
    [
        "Генерация теста",
        "Пишете команду: @methodologist сделай тест...",
        "Читаю шаблон test_template.md, генерирую 10 вопросов, сохраняю .md + .docx",
    ],
    [
        "Контроль качества",
        "Проверяете: режим (4 ч/день), продолжительность (4,5-9 нед), контрольные вопросы, отсутствие выдуманной литературы",
        "Следую жёстким правилам: не выдумываю источники, ставлю контрольные вопросы, соблюдаю режим",
    ],
]

t2 = doc.add_table(rows=len(roles_data), cols=3)
t2.alignment = WD_TABLE_ALIGNMENT.CENTER
t2.style = "Table Grid"
for i, row_data in enumerate(roles_data):
    for j, cell_text in enumerate(row_data):
        cell = t2.cell(i, j)
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(cell_text)
        run.font.size = Pt(9)
        run.font.name = "Times New Roman"
        run.bold = (i == 0)

# ---- 4. ПОРЯДОК ДЕЙСТВИЙ ----
doc.add_heading("4. Порядок действий (пошагово)", level=1)

steps_data = [
    ["Шаг", "Что делать", "Подробно"],
    [
        "1",
        "Подготовить референс",
        "Найдите готовую программу ДПП ПК в .doc/.docx. Я извлеку структуру и создам шаблоны.",
    ],
    [
        "2",
        "Открыть папку в Open Code",
        "Откройте новую сессию Open Code с корнем methodologist/. Агент находится в .opencode/agents/ и будет доступен по @methodologist",
    ],
    [
        "3",
        "Запустить агента (сценарий 1)",
        "Наберите: @methodologist создай программу курса по теме \"...\" на 36 часов, категория — ...",
    ],
    [
        "4",
        "Дождаться результата",
        "Агент сгенерирует .md, затем запустит конвертер → появится .docx. Подтвердите запуск конвертации.",
    ],
    [
        "5",
        "Проверить программу",
        "Откройте .docx в Word. Проверьте: режим 4 ч/день, продолжительность 4,5-9 нед, контрольные вопросы, реальные источники.",
    ],
    [
        "6",
        "Запустить агента (сценарий 2)",
        "Наберите: @methodologist составь план занятия по теме \"...\"",
    ],
    [
        "7",
        "Запустить агента (сценарий 3)",
        "Наберите: @methodologist сделай тест из 10 вопросов по теме \"...\"",
    ],
    [
        "8",
        "Повторить для новой темы",
        "Чтобы создать программу по другой теме — просто напишите новую команду @methodologist. Шаблоны и образцы уже готовы.",
    ],
]

t3 = doc.add_table(rows=len(steps_data), cols=3)
t3.alignment = WD_TABLE_ALIGNMENT.CENTER
t3.style = "Table Grid"
for i, row_data in enumerate(steps_data):
    for j, cell_text in enumerate(row_data):
        cell = t3.cell(i, j)
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(cell_text)
        run.font.size = Pt(10)
        run.font.name = "Times New Roman"
        run.bold = (i == 0)

# ---- 5. ЧТО НУЖНО ОТ ВАС ----
doc.add_heading("5. Какие материалы нужны от вас", level=1)

materials_data = [
    ["Материал", "Зачем", "Формат"],
    ["Референс ДПП ПК", "Анализ структуры, создание шаблона под ваш стандарт", ".doc / .docx"],
    ["Тема программы", "Название курса, который хотите создать", "Текст в команде @methodologist"],
    ["Категория слушателей", "Для кого программа (учителя / СПО / ДОУ / методисты)", "Текст в команде @methodologist"],
    ["Объём в часах", "По умолчанию 36 ч, можно указать другой", "Текст в команде @methodologist"],
    ["Проверка результата", "Вы читаете готовый .docx и говорите, что исправить", "Комментарий в чате"],
]

t4 = doc.add_table(rows=len(materials_data), cols=3)
t4.alignment = WD_TABLE_ALIGNMENT.CENTER
t4.style = "Table Grid"
for i, row_data in enumerate(materials_data):
    for j, cell_text in enumerate(row_data):
        cell = t4.cell(i, j)
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(cell_text)
        run.font.size = Pt(10)
        run.font.name = "Times New Roman"
        run.bold = (i == 0)

# ---- 6. КАК ЗАПУСКАТЬ АГЕНТА ----
doc.add_heading("6. Как запускать агента (шпаргалка)", level=1)

cmds_data = [
    ["Сценарий", "Команда", "Результат"],
    [
        "Программа курса",
        "@methodologist создай программу курса по теме \"...\" на 36 часов, категория — ...",
        ".md + .docx (7 разделов)",
    ],
    [
        "План занятия",
        "@methodologist составь план занятия по теме \"...\"",
        ".md + .docx (хронометраж, активности, ДЗ)",
    ],
    [
        "Тест",
        "@methodologist сделай тест из 10 вопросов по теме \"...\"",
        ".md + .docx (10 вопросов + ключ)",
    ],
]

t5 = doc.add_table(rows=len(cmds_data), cols=3)
t5.alignment = WD_TABLE_ALIGNMENT.CENTER
t5.style = "Table Grid"
for i, row_data in enumerate(cmds_data):
    for j, cell_text in enumerate(row_data):
        cell = t5.cell(i, j)
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(cell_text)
        run.font.size = Pt(10)
        run.font.name = "Times New Roman"
        run.bold = (i == 0)

# ---- 7. ЖЁСТКИЕ ПРАВИЛА ----
doc.add_heading("7. Что агент делает всегда (жёсткие правила)", level=1)

rules = [
    "Режим: 4 часа в день, 1–2 дня в неделю (не 8 часов, не 5 дней подряд)",
    "Продолжительность: 4,5–9 недель (не 1 неделя)",
    "Контрольные вопросы: после каждой темы, минимум 3 вопроса",
    "Литература: только реальные ФЗ, ФГОС, ссылки на официальные сайты. Никаких выдуманных книг.",
    "Практика: ≈ 70% от общего объёма, лекции ≈ 30%",
    "Формат: копирует образец из templates/ — меняет только содержание под новую тему",
    "Конвертация: после .md автоматически создаёт .docx",
]
for r in rules:
    doc.add_paragraph(r, style="List Bullet")

# ---- 8. ИТОГ ----
doc.add_heading("8. Итог проекта", level=1)

summary_data = [
    ["Параметр", "Значение"],
    ["Что создано", "Суб-агент @methodologist (конфиг + шаблоны + образцы + конвертер)"],
    ["Сценариев", "3 (программа, план, тест)"],
    ["Формат на выходе", ".md + .docx (автоконвертация)"],
    ["Режим по умолчанию", "4 ч/день, 1-2 дня/нед, 4,5-9 нед, 36 ч"],
    ["Контрольные вопросы", "Обязательно после каждой темы"],
    ["Литература", "Только реальные источники (ФЗ, ФГОС, ссылки)"],
    ["Коммитов в git", "3 (ядро + конвертер + финальные правки)"],
]

t6 = doc.add_table(rows=len(summary_data), cols=2)
t6.alignment = WD_TABLE_ALIGNMENT.CENTER
t6.style = "Table Grid"
for i, row_data in enumerate(summary_data):
    for j, cell_text in enumerate(row_data):
        cell = t6.cell(i, j)
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(cell_text)
        run.font.size = Pt(10)
        run.font.name = "Times New Roman"
        run.bold = (i == 0)

# ---- SAVE ----
output = (
    r"C:\Users\user\Downloads\for_Cloude Code\prepord-card\methodologist"
    r"\Субагент-Методолог_конспект-инструкция.docx"
)
doc.save(output)
print(f"OK -> {output}")
