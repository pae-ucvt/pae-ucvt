# -*- coding: utf-8 -*-
import xlrd, sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, numbers
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter
from copy import copy

sys.stdout.reconfigure(encoding='utf-8')

# ─── 1. Парсинг исходных файлов ───
files = [
    (r'C:\Users\user\Downloads\for_Cloude Code\prepord-card\Занятие 3_Open Code\Блок 4\численность по площадкам-типам-классам 2-05-24.xls', 2024),
    (r'C:\Users\user\Downloads\for_Cloude Code\prepord-card\Занятие 3_Open Code\Блок 4\площадка-тип_групп-класс_06.05.2025.xls', 2025),
]

def parse_file(filepath):
    wb = xlrd.open_workbook(filepath)
    ws = wb.sheet_by_index(0)
    header_row = None
    for r in range(ws.nrows):
        if ws.cell_value(r, 0) == 'Площадки':
            header_row = r
            break
    detail = []
    current_platform = None
    current_type = None
    for r in range(header_row + 1, ws.nrows):
        col0 = str(ws.cell_value(r, 0)).strip()
        col1 = str(ws.cell_value(r, 1)).strip()
        col2 = str(ws.cell_value(r, 2)).strip()
        col3 = ws.cell_value(r, 3)
        if col3 == '' or col3 == 0:
            continue
        if 'Итого' in col0 or 'Итого' in col2:
            continue
        val = int(col3) if isinstance(col3, float) else col3
        if col0 != '' and col1 == '' and col2 == '':
            current_platform = col0
            current_type = None
        elif col0 == '' and col1 != '' and col2 == '':
            current_type = col1
        elif col0 == '' and col1 == '' and col2 != '':
            cls = str(int(col2)) if isinstance(col2, float) else str(col2)
            detail.append({'platform': current_platform, 'type': current_type, 'class': cls, 'value': int(val)})
    return detail

detail_2024 = parse_file(files[0][0])
detail_2025 = parse_file(files[1][0])
total_24 = sum(d['value'] for d in detail_2024)
total_25 = sum(d['value'] for d in detail_2025)

# Агрегаты
def agg(seq, key):
    d = {}
    for item in seq:
        k = item[key]
        d[k] = d.get(k, 0) + item['value']
    return d

plat_24 = agg(detail_2024, 'platform')
plat_25 = agg(detail_2025, 'platform')
type_24 = agg(detail_2024, 'type')
type_25 = agg(detail_2025, 'type')
class_24 = agg(detail_2024, 'class')
class_25 = agg(detail_2025, 'class')

all_platforms = sorted(set(list(plat_24.keys()) + list(plat_25.keys())))
all_types = sorted(set(list(type_24.keys()) + list(type_25.keys())))
all_classes = sorted(set(list(class_24.keys()) + list(class_25.keys())), key=lambda x: int(x))

# ─── 2. Стили ───
wb = Workbook()

thin = Side(style='thin')
border_all = Border(left=thin, right=thin, top=thin, bottom=thin)
header_font = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
data_font = Font(name='Calibri', size=11)
data_align_right = Alignment(horizontal='right')
data_align_center = Alignment(horizontal='center')
title_font = Font(name='Calibri', bold=True, size=13)
subtitle_font = Font(name='Calibri', bold=True, size=11)
green_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
red_fill = PatternFill(start_color='FCE4EC', end_color='FCE4EC', fill_type='solid')
new_fill = PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid')
total_fill = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')

def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border_all

def style_data_cell(cell, is_number=False, is_pct=False):
    cell.font = data_font
    cell.border = border_all
    if is_number:
        cell.alignment = data_align_right
    elif is_pct:
        cell.alignment = data_align_right

def auto_width(ws, min_w=10, max_w=20):
    for col in ws.columns:
        mx = min_w
        for cell in col:
            if cell.value:
                mx = max(mx, min(len(str(cell.value)) + 2, max_w))
        ws.column_dimensions[get_column_letter(col[0].column)].width = mx

# ─── 3. Лист 1: Площадки ───
ws1 = wb.active
ws1.title = 'Площадки'
ws1.cell(row=1, column=1, value='Сводная таблица по площадкам').font = title_font
ws1.merge_cells('A1:F1')

headers = ['Площадка', '02.05.2024', '06.05.2025', 'Динамика, чел.', 'Δ, %', 'Примечание']
row = 3
for c, h in enumerate(headers, 1):
    ws1.cell(row=row, column=c, value=h)
style_header(ws1, row, len(headers))

for i, p in enumerate(all_platforms, row + 1):
    v24 = plat_24.get(p, 0)
    v25 = plat_25.get(p, 0)
    if v24 == 0:
        dyn = v25
        pct = '—'
        note = 'новая'
    elif v25 == 0:
        dyn = -v24
        pct = '-100,00 %'
        note = 'исчезла'
    else:
        dyn = v25 - v24
        pct_val = (v25 - v24) / v24 * 100
        pct = f'{pct_val:+.2f} %'
        note = ''
    ws1.cell(row=i, column=1, value=p)
    ws1.cell(row=i, column=2, value=v24)
    ws1.cell(row=i, column=3, value=v25)
    ws1.cell(row=i, column=4, value=dyn)
    ws1.cell(row=i, column=5, value=pct)
    ws1.cell(row=i, column=6, value=note)
    for c in range(1, 7):
        style_data_cell(ws1.cell(row=i, column=c), is_number=(c in (2,3,4)))
    if note == 'новая':
        for c in range(1, 7):
            ws1.cell(row=i, column=c).fill = new_fill
    elif isinstance(dyn, (int, float)) and dyn < 0:
        for c in range(1, 7):
            ws1.cell(row=i, column=c).fill = red_fill
    elif isinstance(dyn, (int, float)) and dyn > 0:
        for c in range(1, 7):
            ws1.cell(row=i, column=c).fill = green_fill

# Итого
i += 1
diff_t = total_25 - total_24
pct_t = diff_t / total_24 * 100
ws1.cell(row=i, column=1, value='ИТОГО')
ws1.cell(row=i, column=2, value=total_24)
ws1.cell(row=i, column=3, value=total_25)
ws1.cell(row=i, column=4, value=diff_t)
ws1.cell(row=i, column=5, value=f'{pct_t:+.2f} %')
ws1.cell(row=i, column=6, value='')
for c in range(1, 7):
    cell = ws1.cell(row=i, column=c)
    cell.font = Font(name='Calibri', bold=True, size=11)
    cell.border = border_all
    cell.fill = total_fill
    if c in (2,3,4):
        cell.alignment = data_align_right

auto_width(ws1)

# График по площадкам
last_data = i - 1
chart1 = BarChart()
chart1.type = 'col'
chart1.title = 'Численность по площадкам (2024 vs 2025)'
chart1.y_axis.title = 'Человек'
chart1.x_axis.title = 'Площадка'
chart1.style = 10
chart1.width = 22
chart1.height = 14

cats = Reference(ws1, min_col=1, min_row=4, max_row=last_data)
data_cols = Reference(ws1, min_col=2, max_col=3, min_row=3, max_row=last_data)
chart1.add_data(data_cols, titles_from_data=True)
chart1.set_categories(cats)
chart1.shape = 4

# Назначить цвета
from openpyxl.chart.series import DataPoint
from openpyxl.drawing.fill import PatternFillProperties, ColorChoice
s1 = chart1.series[0]
s1.graphicalProperties.solidFill = '4472C4'
s2 = chart1.series[1]
s2.graphicalProperties.solidFill = 'ED7D31'

ws1.add_chart(chart1, 'A' + str(i + 2))

# ─── 4. Лист 2: Типы групп ───
ws2 = wb.create_sheet('Типы групп')
ws2.cell(row=1, column=1, value='Сводная таблица по типам групп').font = title_font
ws2.merge_cells('A1:F1')

headers2 = ['Тип группы', '02.05.2024', '06.05.2025', 'Динамика, чел.', 'Δ, %', 'Примечание']
row2 = 3
for c, h in enumerate(headers2, 1):
    ws2.cell(row=row2, column=c, value=h)
style_header(ws2, row2, len(headers2))

for i, t in enumerate(all_types, row2 + 1):
    v24 = type_24.get(t, 0)
    v25 = type_25.get(t, 0)
    if v24 == 0:
        dyn = v25
        pct = '—'
        note = 'новый'
    elif v25 == 0:
        dyn = -v24
        pct = '-100,00 %'
        note = 'исчез'
    else:
        dyn = v25 - v24
        pct_val = (v25 - v24) / v24 * 100
        pct = f'{pct_val:+.2f} %'
        note = ''
    ws2.cell(row=i, column=1, value=t)
    ws2.cell(row=i, column=2, value=v24)
    ws2.cell(row=i, column=3, value=v25)
    ws2.cell(row=i, column=4, value=dyn)
    ws2.cell(row=i, column=5, value=pct)
    ws2.cell(row=i, column=6, value=note)
    for c in range(1, 7):
        style_data_cell(ws2.cell(row=i, column=c), is_number=(c in (2,3,4)))
    if note == 'новый':
        for c in range(1, 7):
            ws2.cell(row=i, column=c).fill = new_fill
    elif note == 'исчез':
        for c in range(1, 7):
            ws2.cell(row=i, column=c).fill = PatternFill(start_color='FCE4EC', end_color='FCE4EC', fill_type='solid')
    elif isinstance(dyn, (int, float)) and dyn < 0:
        for c in range(1, 7):
            ws2.cell(row=i, column=c).fill = PatternFill(start_color='FFF3E0', end_color='FFF3E0', fill_type='solid')
    elif isinstance(dyn, (int, float)) and dyn > 0:
        for c in range(1, 7):
            ws2.cell(row=i, column=c).fill = green_fill

i += 1
ws2.cell(row=i, column=1, value='ИТОГО')
ws2.cell(row=i, column=2, value=total_24)
ws2.cell(row=i, column=3, value=total_25)
ws2.cell(row=i, column=4, value=diff_t)
ws2.cell(row=i, column=5, value=f'{pct_t:+.2f} %')
ws2.cell(row=i, column=6, value='')
for c in range(1, 7):
    cell = ws2.cell(row=i, column=c)
    cell.font = Font(name='Calibri', bold=True, size=11)
    cell.border = border_all
    cell.fill = total_fill
    if c in (2,3,4):
        cell.alignment = data_align_right

auto_width(ws2)

# График по типам групп
last_data2 = i - 1
chart2 = BarChart()
chart2.type = 'col'
chart2.title = 'Численность по типам групп (2024 vs 2025)'
chart2.y_axis.title = 'Человек'
chart2.x_axis.title = 'Тип группы'
chart2.style = 10
chart2.width = 26
chart2.height = 14

cats2 = Reference(ws2, min_col=1, min_row=4, max_row=last_data2)
data_cols2 = Reference(ws2, min_col=2, max_col=3, min_row=3, max_row=last_data2)
chart2.add_data(data_cols2, titles_from_data=True)
chart2.set_categories(cats2)
s1 = chart2.series[0]
s1.graphicalProperties.solidFill = '4472C4'
s2 = chart2.series[1]
s2.graphicalProperties.solidFill = 'ED7D31'

ws2.add_chart(chart2, 'A' + str(i + 2))

# ─── 5. Лист 3: Возрастной состав ───
ws3 = wb.create_sheet('Возрастной состав')
ws3.cell(row=1, column=1, value='Распределение по классам').font = title_font
ws3.merge_cells('A1:H1')

headers3 = ['Класс', '02.05.2024', '% 2024', '06.05.2025', '% 2025', 'Δ, чел.', 'Δ, п.п.']
row3 = 3
for c, h in enumerate(headers3, 1):
    ws3.cell(row=row3, column=c, value=h)
style_header(ws3, row3, len(headers3))

for i, cl in enumerate(all_classes, row3 + 1):
    v24 = class_24.get(cl, 0)
    v25 = class_25.get(cl, 0)
    p24 = v24 / total_24 * 100
    p25 = v25 / total_25 * 100
    diff = v25 - v24
    pp = p25 - p24
    ws3.cell(row=i, column=1, value=int(cl))
    ws3.cell(row=i, column=2, value=v24)
    ws3.cell(row=i, column=3, value=f'{p24:.2f} %')
    ws3.cell(row=i, column=4, value=v25)
    ws3.cell(row=i, column=5, value=f'{p25:.2f} %')
    ws3.cell(row=i, column=6, value=diff)
    ws3.cell(row=i, column=7, value=f'{pp:+.2f} п.п.')
    for c in range(1, 8):
        style_data_cell(ws3.cell(row=i, column=c), is_number=(c in (2,4,6)))
    if diff < 0:
        for c in range(1, 8):
            ws3.cell(row=i, column=c).fill = PatternFill(start_color='FFF3E0', end_color='FFF3E0', fill_type='solid')
    elif diff > 0:
        for c in range(1, 8):
            ws3.cell(row=i, column=c).fill = green_fill

i += 1
ws3.cell(row=i, column=1, value='ИТОГО')
ws3.cell(row=i, column=2, value=total_24)
ws3.cell(row=i, column=3, value='100,00 %')
ws3.cell(row=i, column=4, value=total_25)
ws3.cell(row=i, column=5, value='100,00 %')
ws3.cell(row=i, column=6, value=diff_t)
ws3.cell(row=i, column=7, value='')
for c in range(1, 8):
    cell = ws3.cell(row=i, column=c)
    cell.font = Font(name='Calibri', bold=True, size=11)
    cell.border = border_all
    cell.fill = total_fill
    if c in (2,4,6):
        cell.alignment = data_align_right

# Группировка по возрастам
i += 2
ws3.cell(row=i, column=1, value='Группировка по возрастным категориям').font = subtitle_font
ws3.merge_cells(f'A{i}:E{i}')
i += 1
age_headers = ['Категория', '02.05.2024', '% 2024', '06.05.2025', '% 2025']
for c, h in enumerate(age_headers, 1):
    ws3.cell(row=i, column=c, value=h)
style_header(ws3, i, len(age_headers))

age_groups = {
    '1–6 (младшие)': [c for c in all_classes if int(c) <= 6],
    '7–9 (средние)': [c for c in all_classes if 7 <= int(c) <= 9],
    '10–12 (старшие)': [c for c in all_classes if int(c) >= 10],
}

for grp_name, classes in age_groups.items():
    i += 1
    v24 = sum(class_24.get(c, 0) for c in classes)
    v25 = sum(class_25.get(c, 0) for c in classes)
    p24 = v24 / total_24 * 100
    p25 = v25 / total_25 * 100
    ws3.cell(row=i, column=1, value=grp_name)
    ws3.cell(row=i, column=2, value=v24)
    ws3.cell(row=i, column=3, value=f'{p24:.2f} %')
    ws3.cell(row=i, column=4, value=v25)
    ws3.cell(row=i, column=5, value=f'{p25:.2f} %')
    for c in range(1, 6):
        style_data_cell(ws3.cell(row=i, column=c), is_number=(c in (2,4)))

auto_width(ws3)

# График возрастного состава (столбцы 2024 и 2025)
last_row_class = row3 + len(all_classes)
chart3 = BarChart()
chart3.type = 'col'
chart3.title = 'Распределение по классам (2024 vs 2025)'
chart3.y_axis.title = 'Человек'
chart3.x_axis.title = 'Класс'
chart3.style = 10
chart3.width = 22
chart3.height = 14

cats3 = Reference(ws3, min_col=1, min_row=row3 + 1, max_row=last_row_class)
data3a = Reference(ws3, min_col=2, min_row=row3, max_row=last_row_class)
data3b = Reference(ws3, min_col=4, min_row=row3, max_row=last_row_class)
chart3.add_data(data3a, titles_from_data=True)
chart3.add_data(data3b, titles_from_data=True)
chart3.set_categories(cats3)
chart3.series[0].graphicalProperties.solidFill = '4472C4'
chart3.series[1].graphicalProperties.solidFill = 'ED7D31'

ws3.add_chart(chart3, 'A' + str(i + 2))

# ─── 6. Лист 4: Выводы ───
ws4 = wb.create_sheet('Выводы')
ws4.cell(row=1, column=1, value='Аналитические выводы').font = title_font
ws4.merge_cells('A1:D1')

conclusions = [
    ['№', 'Вывод', 'Факт', 'Цифра'],
    [1, 'Общее снижение численности',
     'За год число слушателей сократилось на 59 чел. Падение на всех старых площадках.',
     f'{total_24} → {total_25} ({pct_t:.2f} %)'],
    [2, 'Ушинского и Сизова — главные точки падения',
     'Эти две площадки дали 83 % общего снижения. Ушинского −38,89 %, Сизова −20,79 %.',
     f'Ушинского: 72→44 (−28), Сизова: 101→80 (−21)'],
    [3, 'Исчезновение/появление типов групп',
     'Ушли 3D (16), ВМ (33), ПМ2 (21). Пришли ДИ (9), КМ (27). Чистая потеря −34 чел.',
     'Возможно переименование программ — нужна сверка с методистами'],
    [4, 'Смещение возрастного состава в старшие',
     'Доля младших (1–6) снизилась с 52 % до 48 %. Старших (10–12) — выросла с 7 % до 9 %.',
     'Младших −44 чел., средних+старших −15 чел.'],
    [5, 'Рост ПТ, МП, ОП на фоне общего падения',
     'При общем снижении −16 % группы ПТ (+14 %), МП (+56 %), ОП (+24 %) растут.',
     'ПТ +8, МП +5, ОП +6 — эти направления стоит поддержать'],
]

for r_idx, row_data in enumerate(conclusions, 3):
    for c_idx, val in enumerate(row_data, 1):
        cell = ws4.cell(row=r_idx, column=c_idx, value=val)
        cell.border = border_all
        if r_idx == 3:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
        else:
            cell.font = data_font
            cell.alignment = Alignment(wrap_text=True, vertical='top')

ws4.column_dimensions['A'].width = 5
ws4.column_dimensions['B'].width = 35
ws4.column_dimensions['C'].width = 60
ws4.column_dimensions['D'].width = 45

# ─── Сохраняем ───
out_path = r'C:\Users\user\Downloads\for_Cloude Code\prepord-card\Занятие 3_Open Code\Блок 4\Анализ_2024_vs_2025.xlsx'
wb.save(out_path)
print(f'Готово: {out_path}')
