import os
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

base = r'C:\Users\user\Downloads\for_Cloude Code\prepord-card\Программа ПК - ИИ для учителей'
md_path = os.path.join(base, 'Программа_ПК_ИИ_для_учителей_36ч.md')

with open(md_path, 'r', encoding='utf-8') as f:
    text = f.read()

doc = Document()

section = doc.sections[0]
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.paragraph_format.line_spacing = 1.15
style.paragraph_format.space_after = Pt(4)

for level in [1, 2, 3]:
    hs = doc.styles[f'Heading {level}']
    hs.font.name = 'Calibri'
    hs.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)
    if level == 1:
        hs.font.size = Pt(16)
        hs.font.bold = True
    elif level == 2:
        hs.font.size = Pt(14)
        hs.font.bold = True
    elif level == 3:
        hs.font.size = Pt(12)
        hs.font.bold = True

def add_colored_table(doc, data):
    if not data:
        return
    rows_n = len(data)
    cols_n = len(data[0])
    table = doc.add_table(rows=rows_n, cols=cols_n)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Light Grid Accent 1'
    for ri, row_data in enumerate(data):
        for ci, val in enumerate(row_data):
            cell = table.rows[ri].cells[ci]
            cell.text = str(val)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(9)
                    if ri == 0:
                        run.bold = True

def add_code_block(doc, code_text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), 'F0F0F0')
    shading.set(qn('w:val'), 'clear')
    p.paragraph_format.element.get_or_add_pPr().append(shading)
    run = p.add_run(code_text)
    run.font.size = Pt(8.5)
    run.font.name = 'Courier New'

# ---------- COVER PAGE ----------
for _ in range(6):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('ПРОГРАММА')
run.bold = True
run.font.size = Pt(28)
run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('повышения квалификации')
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('«Введение в искусственный интеллект:\nметодика преподавания для учителей-предметников»')
run.bold = True
run.font.size = Pt(18)
run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

doc.add_paragraph()
doc.add_paragraph()

details = [
    ('Объём', '36 академических часов'),
    ('Категория слушателей', 'Учителя-предметники 5–9 классов'),
    ('Форма обучения', 'Очно-заочная с применением ДОТ'),
    ('Итоговая аттестация', 'Тестирование + защита практической работы'),
]
for label, val in details:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f'{label}: ')
    run.bold = True
    run.font.size = Pt(12)
    run = p.add_run(val)
    run.font.size = Pt(12)

doc.add_page_break()

# ---------- MAIN CONTENT ----------
lines = text.split('\n')
i = 0
in_code = False
code_buf = []
in_yaml = False

while i < len(lines):
    line = lines[i]
    stripped = line.strip()

    if stripped == '---' and not in_yaml and i < 5:
        in_yaml = True
        i += 1
        continue
    if in_yaml:
        if stripped == '---':
            in_yaml = False
        i += 1
        continue

    if stripped.startswith('```'):
        if not in_code:
            in_code = True
            code_buf = []
        else:
            in_code = False
            add_code_block(doc, '\n'.join(code_buf))
            code_buf = []
        i += 1
        continue

    if in_code:
        code_buf.append(line)
        i += 1
        continue

    if stripped == '':
        i += 1
        continue

    if line.startswith('# ') or line.startswith('## ') or line.startswith('### '):
        level = line.count('#')
        txt = line.strip('# ')
        doc.add_heading(txt, level=level)
        i += 1
        continue

    if stripped == '---':
        i += 1
        continue

    if stripped.startswith('|') and '|' in stripped[1:]:
        rows_data = []
        while i < len(lines) and lines[i].strip().startswith('|') and '|' in lines[i].strip()[1:]:
            cols = [c.strip() for c in lines[i].split('|')[1:-1]]
            rows_data.append(cols)
            i += 1
        data_rows = []
        for r in rows_data:
            if not all(c.strip().replace('-', '').replace(':', '') == '' for c in r):
                data_rows.append(r)
        if data_rows and len(data_rows) >= 1:
            add_colored_table(doc, data_rows)
        continue

    if stripped.startswith('**') and stripped.endswith('**') and len(stripped) > 4:
        p = doc.add_paragraph()
        run = p.add_run(stripped.strip('*'))
        run.bold = True
        i += 1
        continue

    if stripped.startswith('- ') or stripped.startswith('* '):
        txt = stripped[2:]
        doc.add_paragraph(txt, style='List Bullet')
        i += 1
        continue

    match_nl = re.match(r'^(\d+)\.\s*(.*)', stripped)
    if match_nl:
        doc.add_paragraph(match_nl.group(2), style='List Number')
        i += 1
        continue

    p = doc.add_paragraph()
    parts = re.split(r'(\*\*[^*]+\*\*)', stripped)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            run = p.add_run(part)
        run.font.size = Pt(11)
    i += 1

for section in doc.sections:
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Страница ')
    run.font.size = Pt(9)
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    run2 = p.add_run()
    run2._r.append(fld_char1)
    run3 = p.add_run()
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE '
    run3._r.append(instr)
    run4 = p.add_run()
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run4._r.append(fld_char2)
    for r in [run, run2, run3, run4]:
        r.font.size = Pt(9)

docx_path = os.path.join(base, 'Программа_ПК_ИИ_для_учителей_36ч.docx')
doc.save(docx_path)
print(f'DOCX created: {docx_path}')
print(f'Size: {os.path.getsize(docx_path)} B')
