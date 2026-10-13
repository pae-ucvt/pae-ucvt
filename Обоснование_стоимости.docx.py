from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

doc = Document()

section = doc.sections[0]
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.paragraph_format.line_spacing = 1.15
style.paragraph_format.space_after = Pt(4)

def add_colored_table(doc, data, col_widths=None):
    if not data:
        return
    rows_n = len(data)
    cols_n = len(data[0])
    table = doc.add_table(rows=rows_n, cols=cols_n)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row_data in enumerate(data):
        for ci, val in enumerate(row_data):
            cell = table.rows[ri].cells[ci]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(9)
            if ri == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), '1A3C6E')
                shading.set(qn('w:val'), 'clear')
                cell._element.get_or_add_tcPr().append(shading)
            elif ri % 2 == 0:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), 'F2F5FA')
                shading.set(qn('w:val'), 'clear')
                cell._element.get_or_add_tcPr().append(shading)
    if col_widths:
        for row in table.rows:
            for ci, w in enumerate(col_widths):
                if ci < len(row.cells):
                    row.cells[ci].width = Cm(w)
    return table


# ===== COVER =====
for _ in range(4):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('ОБОСНОВАНИЕ СТОИМОСТИ УСЛУГ')
run.bold = True
run.font.size = Pt(18)
run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('по программе\n«Основы технологий искусственного интеллекта» (36 ч)')
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.add_paragraph()
doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Разработчик программы:')
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('[Ваше Имя]')
run.bold = True
run.font.size = Pt(13)

doc.add_page_break()

# ===== 1. ОБЩАЯ ИНФОРМАЦИЯ =====
p = doc.add_paragraph()
run = p.add_run('1. Общая информация')
run.bold = True
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

doc.add_paragraph()

p = doc.add_paragraph()
run = p.add_run(
    'Настоящий документ содержит обоснование стоимости услуг по разработке'
    ' и проведению программы повышения квалификации'
    ' «Основы технологий искусственного интеллекта» (36 академических часов, 42 слушателя, 3 группы).'
)
run.font.size = Pt(11)

p = doc.add_paragraph()
run = p.add_run(
    'Программа разрабатывается с нуля под конкретные требования Заказчика,'
    ' не является типовой или адаптированной из существующих курсов.'
    ' Все материалы — конспекты, презентации, раздаточные материалы, тесты,'
    ' сценарии вебинаров, инструкции для СДО — создаются впервые и являются'
    ' авторской разработкой Исполнителя.'
)
run.font.size = Pt(11)

doc.add_paragraph()

# ===== 2. ПЕРЕЧЕНЬ РАБОТ =====
p = doc.add_paragraph()
run = p.add_run('2. Пошаговый перечень работ по разработке программы')
run.bold = True
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

doc.add_paragraph()

# 2.1 Учебный план
p = doc.add_paragraph()
run = p.add_run('2.1. Разработка учебного плана и рабочей программы')
run.bold = True
run.font.size = Pt(11)

p = doc.add_paragraph()
run = p.add_run(
    'Составление детального учебно-тематического плана на 36 часов:'
    ' распределение часов по модулям и темам, определение форм занятий'
    ' (лекции, практика, самостоятельная работа, вебинары),'
    ' формулировка планируемых результатов. Оформление в соответствии'
    ' с требованиями НОУР и Приложением № 2 к ТЗ.'
)
run.font.size = Pt(11)

doc.add_paragraph()

# 2.2 Модули
p = doc.add_paragraph()
run = p.add_run('2.2. Разработка содержания по модулям')
run.bold = True
run.font.size = Pt(11)

doc.add_paragraph()

# Table: modules breakdown
modules_header = [
    '№', 'Модуль / тема', 'Конспект\n(ч)', 'Презентация\n(ч)',
    'Раздатка\n(ч)', 'Практика /\nзадания (ч)', 'Всего\n(ч)'
]
modules_data = [modules_header]

module_rows = [
    ('1', 'Основы систем ИИ для непрофильных специалистов', '5', '3', '2', '3', '13'),
    ('2', 'ИИ в профессиональной деятельности: возможности,\nограничения, сферы применения', '4', '2', '1', '2', '9'),
    ('3', 'Большие языковые модели: генерация, анализ\nтекста, работа с документами', '5', '3', '2', '3', '13'),
    ('4', 'RAG-технологии: интеллектуальный поиск,\nкорпоративные базы знаний', '4', '2', '2', '2', '10'),
    ('5', 'Компьютерное зрение в задачах управления\nи автоматизации', '4', '2', '1', '2', '9'),
    ('6', 'Правовые и организационные аспекты применения ИИ', '4', '2', '1', '1', '8'),
    ('7', 'Организация профессиональной деятельности\nс использованием инструментов ИИ', '4', '2', '1', '2', '9'),
]

for row in module_rows:
    modules_data.append(list(row))

modules_data.append(['', 'Итого по модулям', '30', '16', '10', '15', '71'])

add_colored_table(doc, modules_data, col_widths=[0.8, 6.5, 1.3, 1.3, 1.0, 1.3, 1.0])

doc.add_paragraph()

# 2.3 Тесты
p = doc.add_paragraph()
run = p.add_run('2.3. Разработка тестовых материалов')
run.bold = True
run.font.size = Pt(11)

test_data = [
    ['Вид работы', 'Описание', 'Часы'],
    ['Итоговый тест', '15 вопросов × 4 варианта ответа, 1 верный.\nРазработка, проверка, оформление.', '8'],
    ['Тесты с ответами\nдля самопроверки', 'По 5 вопросов к каждому модулю\n(всего 35 вопросов + ключи).', '10'],
    ['Промежуточные тесты\n(по желанию заказчика)', '3 теста по 10 вопросов\nс ответами для самопроверки.', '6'],
    ['', 'Итого по тестам', '18–24'],
]
add_colored_table(doc, test_data, col_widths=[3.5, 6.0, 1.5])

doc.add_paragraph()

# 2.4 Вебинары
p = doc.add_paragraph()
run = p.add_run('2.4. Разработка сценариев вебинаров (3 шт.)')
run.bold = True
run.font.size = Pt(11)

web_data = [
    ['Вид работы', 'Описание', 'Часы'],
    ['Сценарий вебинара №1', 'Введение в ИИ, LLM, демонстрация\nинструментов. Структура, тайминг, промпты.', '5'],
    ['Сценарий вебинара №2', 'RAG, компьютерное зрение, кейсы.\nПлан демонстраций и практики.', '5'],
    ['Сценарий вебинара №3', 'Правовые аспекты, организация работы\nс ИИ. Итоговое тестирование.', '4'],
    ['', 'Итого по вебинарам', '14'],
]
add_colored_table(doc, web_data, col_widths=[3.5, 6.0, 1.5])

doc.add_paragraph()

# 2.5 СДО
p = doc.add_paragraph()
run = p.add_run('2.5. Подготовка материалов для системы дистанционного обучения')
run.bold = True
run.font.size = Pt(11)

sdo_data = [
    ['Вид работы', 'Часы'],
    ['Адаптация конспектов и презентаций для загрузки в СДО', '4'],
    ['Подготовка инструкции для слушателей по работе в СДО', '3'],
    ['Создание тестовых заданий в формате СДО', '4'],
    ['Проверка работоспособности материалов, исправление ошибок', '3'],
    ['Итого по СДО', '14'],
]
add_colored_table(doc, sdo_data, col_widths=[7.0, 1.5])

doc.add_paragraph()

# 2.6 Дополнительно
p = doc.add_paragraph()
run = p.add_run('2.6. Дополнительные работы')
run.bold = True
run.font.size = Pt(11)

extra_data = [
    ['Вид работы', 'Часы'],
    ['Согласование программы с заказчиком, внесение правок\n(до 2 итераций)', '6'],
    ['Методическая редактура и вычитка всех материалов', '5'],
    ['Упаковка и передача комплекта материалов Заказчику', '2'],
    ['Итого по дополнительным работам', '13'],
]
add_colored_table(doc, extra_data, col_widths=[7.0, 1.5])

doc.add_paragraph()

# ===== 3. СВОДНЫЙ РАСЧЁТ =====
p = doc.add_paragraph()
run = p.add_run('3. Сводный расчёт трудозатрат на разработку')
run.bold = True
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

doc.add_paragraph()

total_data = [
    ['Этап работ', 'Часы'],
    ['Учебный план и рабочая программа', '8'],
    ['Конспекты по 7 модулям', '30'],
    ['Презентации по 7 модулям', '16'],
    ['Раздаточные материалы по 7 модулям', '10'],
    ['Практические задания по 7 модулям', '15'],
    ['Тестовые материалы (итоговый + самопроверка)', '24'],
    ['Сценарии вебинаров (3 шт.)', '14'],
    ['Подготовка материалов для СДО', '14'],
    ['Согласование, доработка, редактура', '13'],
    ['Итого разработка', '~144'],
]
add_colored_table(doc, total_data, col_widths=[8.0, 2.0])

doc.add_paragraph()

p = doc.add_paragraph()
run = p.add_run(
    'Общий объём работ по разработке: ориентировочно 144 часа.'
    '\nДанная программа создаётся с нуля: не используется готовая типовая программа,'
    ' не применяются шаблоны прошлых курсов. Каждый модуль, каждый слайд,'
    ' каждое тестовое задание разрабатывается индивидуально под требования'
    ' конкретного Заказчика и под целевую аудиторию слушателей.'
)
run.font.size = Pt(11)

doc.add_paragraph()

# ===== 4. СТОИМОСТЬ =====
p = doc.add_paragraph()
run = p.add_run('4. Расчёт стоимости')
run.bold = True
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

doc.add_paragraph()

# 4.1 Teaching
p = doc.add_paragraph()
run = p.add_run('4.1. Преподавание')
run.bold = True
run.font.size = Pt(12)

p = doc.add_paragraph()
run = p.add_run(
    'Стоимость 1 академического часа преподавания — 1 500 руб.'
    ' (стандартная ставка НОУР).'
    ' Расчёт производится по факту проведённых занятий.'
)
run.font.size = Pt(11)

teach_data = [
    ['Параметр', 'Значение'],
    ['Ставка за 1 ак. час', '1 500 руб.'],
    ['Количество групп', '3'],
    ['Кол-во часов на 1 группу (синхронная работа)', 'уточняется'],
    ['Итого преподавание', 'от 54 000 до 162 000 руб.'],
]
add_colored_table(doc, teach_data, col_widths=[4.0, 4.0])

doc.add_paragraph()

# 4.2 Development
p = doc.add_paragraph()
run = p.add_run('4.2. Разработка программы и материалов')
run.bold = True
run.font.size = Pt(12)

p = doc.add_paragraph()
run = p.add_run(
    'Разработка программы не является преподавательской работой.'
    ' Она включает: анализ требований заказчика, проектирование структуры,'
    ' написание авторских текстов, создание визуальных материалов,'
    ' разработку тестового фонда, подготовку сценариев, адаптацию'
    ' под СДО, согласование и доработку.'
    '\n\n'
    'Это интеллектуальная, творческая и методическая работа, которая:'
    '\n  • не может быть выполнена в рамках почасовой преподавательской ставки;'
    '\n  • создаёт готовый продукт, остающийся в распоряжении НОУР'
    ' и пригодный для многократного использования;'
    '\n  • требует высокой квалификации (знание предмета, педагогический дизайн,'
    ' владение инструментами ИИ, опыт разработки учебных программ).'
)
run.font.size = Pt(11)

doc.add_paragraph()

cost_data = [
    ['Параметр', 'Значение'],
    ['Объём работ по разработке', '144 часа'],
    ['Ставка за 1 час разработки (авторская)', '2 000 руб.'],
    ['Итого разработка', '288 000 руб.'],
]
add_colored_table(doc, cost_data, col_widths=[5.0, 4.0])

doc.add_paragraph()

# 4.3 Total
p = doc.add_paragraph()
run = p.add_run('4.3. Итоговая стоимость услуг')
run.bold = True
run.font.size = Pt(12)

doc.add_paragraph()

total_cost_data = [
    ['Составляющая', 'Сумма'],
    ['Разработка программы и всех материалов (144 ч × 2 000 руб./ч)', '288 000 руб.'],
    ['Преподавание (3 группы, часы уточняются)', 'от 54 000 до 162 000 руб.'],
    ['Итого (ориентировочно)', 'от 342 000 до 450 000 руб.'],
]
add_colored_table(doc, total_cost_data, col_widths=[6.0, 4.0])

doc.add_paragraph()

# ===== 5. ЗАКЛЮЧЕНИЕ =====
p = doc.add_paragraph()
run = p.add_run('5. Заключение')
run.bold = True
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

doc.add_paragraph()

p = doc.add_paragraph()
run = p.add_run(
    'Прошу утвердить стоимость разработки программы в размере 288 000 руб.'
    ' (двести восемьдесят восемь тысяч рублей) из расчёта 144 часов × 2 000 руб./ч.'
    '\n\n'
    'Стоимость преподавания определяется дополнительно по факту'
    ' согласования количества синхронных часов на одну группу.'
    '\n\n'
    'Готова предоставить детализацию по любому из пунктов и'
    ' приступить к разработке после утверждения стоимости.'
)
run.font.size = Pt(11)

doc.add_paragraph()
doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = p.add_run('С уважением,')
run.font.size = Pt(11)

p = doc.add_paragraph()
run = p.add_run('[Ваше Имя]')
run.bold = True
run.font.size = Pt(12)

p = doc.add_paragraph()
run = p.add_run('[Дата]')
run.font.size = Pt(11)

# ---- Footer ----
for section in doc.sections:
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Обоснование стоимости — «Основы технологий искусственного интеллекта» (36 ч)')
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

# ---- Save ----
output_path = r'C:\Users\user\Downloads\for_Cloude Code\prepord-card\Обоснование_стоимости.docx'
doc.save(output_path)
print(f'OK: {output_path}')
print(f'Size: {os.path.getsize(output_path)} B')
