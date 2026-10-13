# -*- coding: utf-8 -*-
import xlrd, sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

# ─── 1. Парсинг исходных файлов ───
files = [
    (r'C:\Users\user\Downloads\for_Cloude Code\prepord-card\Занятие 3_Open Code\Блок 4\численность по площадкам-типам-классам 2-05-24.xls', 2024),
    (r'C:\Users\user\Downloads\for_Cloude Code\prepord-card\Занятие 3_Open Code\Блок 4\площадка-тип_групп-класс_06.05.2025.xls', 2025),
]

def parse_file(filepath):
    wb = xlrd.open_workbook(filepath)
    ws = wb.sheet_by_index(0)
    header_row = None
    for r in range(ws.nrows):
        if ws.cell_value(r, 0) == 'Площадки':
            header_row = r
            break
    detail = []
    current_platform = None
    current_type = None
    for r in range(header_row + 1, ws.nrows):
        col0 = str(ws.cell_value(r, 0)).strip()
        col1 = str(ws.cell_value(r, 1)).strip()
        col2 = str(ws.cell_value(r, 2)).strip()
        col3 = ws.cell_value(r, 3)
        if col3 == '' or col3 == 0:
            continue
        if 'Итого' in col0 or 'Итого' in col2:
            continue
        val = int(col3) if isinstance(col3, float) else col3
        if col0 != '' and col1 == '' and col2 == '':
            current_platform = col0
            current_type = None
        elif col0 == '' and col1 != '' and col2 == '':
            current_type = col1
        elif col0 == '' and col1 == '' and col2 != '':
            cls = str(int(col2)) if isinstance(col2, float) else str(col2)
            detail.append({'platform': current_platform, 'type': current_type, 'class': cls, 'value': int(val)})
    return detail

d24 = parse_file(files[0][0])
d25 = parse_file(files[1][0])
total_24 = sum(d['value'] for d in d24)
total_25 = sum(d['value'] for d in d25)

def agg(seq, key):
    d = {}
    for item in seq:
        k = item[key]
        d[k] = d.get(k, 0) + item['value']
    return d

p24 = agg(d24, 'platform')
p25 = agg(d25, 'platform')
t24 = agg(d24, 'type')
t25 = agg(d25, 'type')
c24 = agg(d24, 'class')
c25 = agg(d25, 'class')

all_platforms = sorted(set(list(p24.keys()) + list(p25.keys())))
all_types = sorted(set(list(t24.keys()) + list(t25.keys())))
all_classes = sorted(set(list(c24.keys()) + list(c25.keys())), key=lambda x: int(x))

# ─── 2. Стили ───
wb = Workbook()
thin = Side(style='thin')
side_m = Side(style='medium')
b_all = Border(left=thin, right=thin, top=thin, bottom=thin)
b_m = Border(left=side_m, right=side_m, top=side_m, bottom=side_m)
h_font = Font(name='Calibri', bold=True, size=10, color='FFFFFF')
h_fill = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
d_font = Font(name='Calibri', size=10)
db_font = Font(name='Calibri', size=10, bold=True)
t_font = Font(name='Calibri', bold=True, size=16, color='2F5496')
k_big = Font(name='Calibri', bold=True, size=28, color='2F5496')
k_lab = Font(name='Calibri', size=9, color='666666')
k_red = Font(name='Calibri', bold=True, size=16, color='C00000')
k_grn = Font(name='Calibri', bold=True, size=16, color='006100')
s_font = Font(name='Calibri', bold=True, size=12, color='2F5496')
k_bg = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
grn_f = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
red_f = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
new_f = PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid')
tot_f = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')
alt_f = PatternFill(start_color='F2F7FB', end_color='F2F7FB', fill_type='solid')
c_red = Font(name='Calibri', bold=True, size=10, color='C00000')
c_grn = Font(name='Calibri', bold=True, size=10, color='006100')

def center(cell):
    cell.alignment = Alignment(horizontal='center', vertical='center')

def hdr(ws, row, cols):
    for c in range(1, cols + 1):
        cl = ws.cell(row=row+1, column=c) if True else None  # dummy
    for c in range(1, cols + 1):
        cl = ws.cell(row=row, column=c)
        cl.font = h_font; cl.fill = h_fill; cl.border = b_all
        center(cl)

def aw(ws, mn=10, mx=22):
    for col in ws.columns:
        w = mn
        for cl in col:
            if cl.value:
                w = max(w, min(len(str(cl.value)) + 2, mx))
        ws.column_dimensions[get_column_letter(col[0].column)].width = w

# ================================================================
# ЛИСТ 1: Исходные данные
# ================================================================
wsd = wb.active
wsd.title = 'Данные'

def mk_tbl(wsd, srow, scol, items, d1, d2, label):
    wsd.cell(row=srow, column=scol, value=label).font = Font(name='Calibri', bold=True, size=12)
    hdr(wsd, srow + 1, scol + 2)
    headers = ['', '2024', '2025']
    for c, h in enumerate(headers, scol):
        wsd.cell(row=srow + 1, column=c, value=h)
    for i, k in enumerate(items, srow + 2):
        wsd.cell(row=i, column=scol, value=k).font = d_font
        wsd.cell(row=i, column=scol).border = b_all
        for ci, v in [(scol + 1, d1.get(k, 0)), (scol + 2, d2.get(k, 0))]:
            c = wsd.cell(row=i, column=ci, value=v)
            c.font = d_font; c.border = b_all; center(c)
        if i % 2 == 0:
            for cc in range(scol, scol + 3):
                wsd.cell(row=i, column=cc).fill = alt_f

mk_tbl(wsd, 1, 1, all_platforms, p24, p25, 'Площадки')
mk_tbl(wsd, 1, 5, all_types, t24, t25, 'Типы групп')
mk_tbl(wsd, 1, 10, all_classes, c24, c25, 'Распределение по классам')
for r in range(3, 3 + len(all_classes)):
    wsd.cell(row=r, column=10).value = int(wsd.cell(row=r, column=10).value)
aw(wsd)

# ================================================================
# ЛИСТ 2: Дашборд  
# ================================================================
ws = wb.create_sheet('Дашборд')

ws.merge_cells('A1:N1')
c = ws.cell(row=1, column=1, value='ДАШБОРД: Численность слушателей 2024 vs 2025'); c.font = t_font
ws.merge_cells('A2:N2')
c = ws.cell(row=2, column=1, value='Курс «Компьютерные технологии для школьников»')
c.font = Font(name='Calibri', size=10, color='888888')

def kpi(ws, r, cl, cl2, label, val, big_font=k_big, bg=k_bg, bdr=b_m):
    col_start = ws[cl + str(r)].column
    col_end = ws[cl2 + str(r)].column
    ws.merge_cells(f'{cl}{r}:{cl2}{r}')
    c = ws.cell(row=r, column=col_start, value=label)
    c.font = k_lab; center(c)
    for cc in range(col_start, col_end + 1):
        ws.cell(row=r, column=cc).fill = bg
        ws.cell(row=r, column=cc).border = bdr
    ws.merge_cells(f'{cl}{r+1}:{cl2}{r+1}')
    c = ws.cell(row=r+1, column=col_start, value=val)
    c.font = big_font; center(c)
    for cc in range(col_start, col_end + 1):
        ws.cell(row=r+1, column=cc).fill = bg
        ws.cell(row=r+1, column=cc).border = bdr

kpi(ws, 4, 'B', 'C', 'ВСЕГО 2024', total_24)
kpi(ws, 4, 'E', 'F', 'ВСЕГО 2025', total_25)

diff = total_25 - total_24
pct = diff / total_24 * 100
pct_str = f'{pct:+.2f} %'
if diff < 0:
    kpi(ws, 4, 'H', 'I', 'ИЗМЕНЕНИЕ', pct_str, k_red, red_f)
else:
    kpi(ws, 4, 'H', 'I', 'ИЗМЕНЕНИЕ', pct_str, k_grn, grn_f)

kpi(ws, 8, 'B', 'C', 'ПЛОЩАДОК', len(all_platforms))
kpi(ws, 8, 'E', 'F', 'ТИПОВ ГРУПП', len(all_types))

# ─── График 1: Площадки ───
ws.cell(row=12, column=1, value='▍ Численность по площадкам').font = s_font
ch1 = BarChart()
ch1.type = 'col'; ch1.style = 10; ch1.width = 18; ch1.height = 10
ch1.legend.position = 'b'
c1 = Reference(wsd, min_col=1, min_row=3, max_row=2+len(all_platforms))
d1 = Reference(wsd, min_col=2, max_col=3, min_row=2, max_row=2+len(all_platforms))
ch1.add_data(d1, titles_from_data=True); ch1.set_categories(c1)
ch1.series[0].graphicalProperties.solidFill = '4472C4'
ch1.series[1].graphicalProperties.solidFill = 'ED7D31'
for s in ch1.series:
    s.dLbls = DataLabelList(); s.dLbls.showVal = True; s.dLbls.numFmt = '0'
    s.dLbls.font = Font(size=9)
ws.add_chart(ch1, 'A13')

# ─── График 2: Типы групп ───
ws.cell(row=30, column=1, value='▍ Численность по типам групп').font = s_font
ch2 = BarChart()
ch2.type = 'col'; ch2.style = 10; ch2.width = 20; ch2.height = 10
ch2.legend.position = 'b'
c2 = Reference(wsd, min_col=5, min_row=3, max_row=2+len(all_types))
d2 = Reference(wsd, min_col=6, max_col=7, min_row=2, max_row=2+len(all_types))
ch2.add_data(d2, titles_from_data=True); ch2.set_categories(c2)
ch2.series[0].graphicalProperties.solidFill = '4472C4'
ch2.series[1].graphicalProperties.solidFill = 'ED7D31'
for s in ch2.series:
    s.dLbls = DataLabelList(); s.dLbls.showVal = True; s.dLbls.numFmt = '0'
    s.dLbls.font = Font(size=8)
ws.add_chart(ch2, 'A31')

# ─── График 3: Возрастной состав ───
ws.cell(row=48, column=1, value='▍ Распределение по классам').font = s_font
ch3 = BarChart()
ch3.type = 'col'; ch3.style = 10; ch3.width = 24; ch3.height = 10
ch3.legend.position = 'b'
c3 = Reference(wsd, min_col=10, min_row=3, max_row=2+len(all_classes))
d3 = Reference(wsd, min_col=11, max_col=12, min_row=2, max_row=2+len(all_classes))
ch3.add_data(d3, titles_from_data=True); ch3.set_categories(c3)
ch3.series[0].graphicalProperties.solidFill = '4472C4'
ch3.series[1].graphicalProperties.solidFill = 'ED7D31'
for s in ch3.series:
    s.dLbls = DataLabelList(); s.dLbls.showVal = True; s.dLbls.numFmt = '0'
    s.dLbls.font = Font(size=9)
ws.add_chart(ch3, 'A49')

# ─── Таблица: Динамика по площадкам ───
tr = 66
ws.cell(row=tr, column=1, value='▍ Динамика по площадкам').font = s_font
tr += 1
for c, h in enumerate(['Площадка', '2024', '2025', 'Δ, чел.', 'Δ, %'], 1):
    ws.cell(row=tr, column=c, value=h)
hdr(ws, tr, 5)

for i, p in enumerate(all_platforms, tr + 1):
    v24 = p24.get(p, 0); v25 = p25.get(p, 0)
    ws.cell(row=i, column=1, value=p).font = d_font; ws.cell(row=i, column=1).border = b_all
    for ci, v in [(2, v24), (3, v25)]:
        c = ws.cell(row=i, column=ci, value=v); c.font = d_font; center(c); c.border = b_all
    if v24 == 0:
        c4 = ws.cell(row=i, column=4, value=f'+{v25}'); c4.font = d_font; center(c4); c4.border = b_all
        c5 = ws.cell(row=i, column=5, value='новая'); c5.font = d_font; center(c5); c5.border = b_all
        for cc in range(1, 6): ws.cell(row=i, column=cc).fill = new_f
    elif v25 == 0:
        c4 = ws.cell(row=i, column=4, value=f'-{v24}'); c4.font = d_font; center(c4); c4.border = b_all
        c5 = ws.cell(row=i, column=5, value='-100,00 %'); c5.font = d_font; center(c5); c5.border = b_all
    else:
        d = v25 - v24; p = d / v24 * 100
        c4 = ws.cell(row=i, column=4, value=d); c4.font = db_font; center(c4); c4.border = b_all
        c5 = ws.cell(row=i, column=5, value=f'{p:+.2f} %'); c5.font = db_font; center(c5); c5.border = b_all
        if d < 0: c4.font = c_red; c5.font = c_red
        elif d > 0: c4.font = c_grn; c5.font = c_grn
tr2 = tr + 1 + len(all_platforms)
for ci, v in [(1, 'ИТОГО'), (2, total_24), (3, total_25), (4, diff), (5, f'{pct:+.2f} %')]:
    c = ws.cell(row=tr2, column=ci, value=v)
    c.font = db_font; c.fill = tot_f; c.border = b_all; center(c)

# ─── Таблица: Динамика по типам групп ───
tc = 7
ws.cell(row=tr2 + 2, column=tc, value='▍ Динамика по типам групп').font = s_font
for c, h in enumerate(['Тип', '2024', '2025', 'Δ, чел.', 'Δ, %'], tc):
    ws.cell(row=tr2 + 3, column=c, value=h)
hdr(ws, tr2 + 3, tc + 4)

for i, t in enumerate(all_types, tr2 + 4):
    v24 = t24.get(t, 0); v25 = t25.get(t, 0)
    ws.cell(row=i, column=tc, value=t).font = d_font; ws.cell(row=i, column=tc).border = b_all
    for ci, v in [(tc + 1, v24), (tc + 2, v25)]:
        c = ws.cell(row=i, column=ci, value=v); c.font = d_font; center(c); c.border = b_all
    if v24 == 0:
        c4 = ws.cell(row=i, column=tc + 3, value=f'+{v25}'); c4.font = d_font; center(c4); c4.border = b_all
        c5 = ws.cell(row=i, column=tc + 4, value='новый'); c5.font = d_font; center(c5); c5.border = b_all
        for cc in range(tc, tc + 5): ws.cell(row=i, column=cc).fill = new_f
    elif v25 == 0:
        c4 = ws.cell(row=i, column=tc + 3, value=f'-{v24}'); c4.font = d_font; center(c4); c4.border = b_all
        c5 = ws.cell(row=i, column=tc + 4, value='-100,00 %'); c5.font = d_font; center(c5); c5.border = b_all
    else:
        d = v25 - v24; p = d / v24 * 100
        c4 = ws.cell(row=i, column=tc + 3, value=d); c4.font = db_font; center(c4); c4.border = b_all
        c5 = ws.cell(row=i, column=tc + 4, value=f'{p:+.2f} %'); c5.font = db_font; center(c5); c5.border = b_all
        if d < 0: c4.font = c_red; c5.font = c_red
        elif d > 0: c4.font = c_grn; c5.font = c_grn
    for cc in range(tc, tc + 5): ws.cell(row=i, column=cc).border = b_all

tr3 = tr2 + 4 + len(all_types)
for ci, v in [(tc, 'ИТОГО'), (tc + 1, total_24), (tc + 2, total_25), (tc + 3, diff), (tc + 4, f'{pct:+.2f} %')]:
    c = ws.cell(row=tr3, column=ci, value=v)
    c.font = db_font; c.fill = tot_f; c.border = b_all; center(c)

# Ширина столбцов
ws.column_dimensions['A'].width = 2
for l in ['B','C','E','F','H','I']:
    ws.column_dimensions[l].width = 12
for l in ['D','G']:
    ws.column_dimensions[l].width = 3

# ─── Сохраняем ───
out = r'C:\Users\user\Downloads\for_Cloude Code\prepord-card\Занятие 3_Open Code\Блок 4\Дашборд_2024_vs_2025.xlsx'
wb.save(out)
print(f'Готово: {out}')
