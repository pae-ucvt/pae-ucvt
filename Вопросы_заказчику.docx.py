from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

doc = Document()

section = doc.sections[0]
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.paragraph_format.line_spacing = 1.15
style.paragraph_format.space_after = Pt(4)

# ---- Header ----
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = p.add_run('От: ')
run.bold = True
run.font.size = Pt(11)
run = p.add_run('[Ваше имя], преподаватель программы')
run.font.size = Pt(11)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = p.add_run('Тема: ')
run.bold = True
run.font.size = Pt(11)
run = p.add_run('Вопросы для уточнения требований к программе «Основы технологий искусственного интеллекта»')
run.font.size = Pt(11)

doc.add_paragraph()

# ---- Title ----
p = doc.add_paragraph()
run = p.add_run('ВОПРОСЫ ДЛЯ УТОЧНЕНИЯ ТРЕБОВАНИЙ')
run.bold = True
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

p = doc.add_paragraph()
run = p.add_run('к программе «Основы технологий искусственного интеллекта» (36 ч)')
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.add_paragraph()

# ---- Intro ----
p = doc.add_paragraph()
run = p.add_run(
    'Для корректной подготовки программы, учебных материалов и определения'
    ' стоимости услуг прошу уточнить следующие параметры:'
)
run.font.size = Pt(11)

doc.add_paragraph()

# ---- Table ----
questions = [
    ('1',
     'Кто будет учиться?\n'
     'Укажите должности и сферу деятельности слушателей'
     ' (госслужащие, офисные сотрудники, руководители, ИТ-специалисты, смешанная группа).',
     'Определяет сложность примеров, практических кейсов и глубину тем.'
    ),
    ('2',
     'Какая платформа используется для заочной части (Moodle, своя СДО, другая)?\n'
     'Какие форматы контента она поддерживает (SCORM, видео, тесты, PDF)?',
     'Материалы нужно готовить под конкретную систему, чтобы избежать переделок.'
    ),
    ('3',
     'Каков график обучения?\n'
     '36 часов — это интенсив (одна неделя) или распределённый формат (месяц и более)?'
     ' Сколько часов из 36 отводится на вебинары?',
     'Влияет на соотношение теории и практики, объём синхронной и асинхронной работы.'
    ),
    ('4',
     'Что входит в комплект материалов?\n'
     'Есть ли точный перечень (конспект, презентации, раздатка, тесты)?'
     ' Нужна ли рабочая программа (учебный план) по утверждённой форме НОУР?',
     'Чтобы подготовить ровно то, что требуется, без лишней работы.'
    ),
    ('5',
     'Нужна ли очная часть?\n'
     'Если да — сколько часов очно, сколько заочно?',
     'Определяет формат вебинаров и долю самостоятельной работы.'
    ),
    ('6',
     'В каком формате сдавать материалы? (Word, PDF, PPT, вёрстка в СДО)\n'
     'Есть ли корпоративный стиль НОУР (шаблоны, логотип, шрифты, цвета)?',
     'Чтобы сдать в приёмлемом виде с первого раза.'
    ),
    ('7',
     'Какие российские ИИ-сервисы показывать слушателям?\n'
     'Только бесплатные или у НОУР есть доступ к платным версиям'
     ' (GigaChat, YandexGPT, Kandinsky и др.)?',
     'От этого зависит, что слушатели смогут попробовать на практике.'
    ),
    ('8',
     'Тест 15 вопросов — это только итоговый?\n'
     'Или нужны ещё промежуточные тесты по модулям?',
     'Чтобы заложить объём работы по тестовым материалам.'
    ),
    ('9',
     'Какие требования к оформлению удостоверений и отчётных документов?\n'
     'Что нужно предоставить для успешного закрытия договора (акты, списки, скриншоты)?',
     'Чтобы процесс сдачи прошёл гладко.'
    ),
    ('10',
     'Приложение 2 к ТЗ — это содержание программы?\n'
     'Можно получить текст приложения для точного соответствия требованиям?',
     'Чтобы программа совпадала с ожиданиями заказчика.'
    ),
    ('11',
     'Сколько часов синхронной работы с преподавателем предусмотрено на одну группу?\n'
     'Все 36 ч программы — это время с преподавателем, или часть отводится на самостоятельную работу слушателей?',
     'От этого зависит расчёт стоимости преподавательской работы.'
    ),
]

table = doc.add_table(rows=len(questions) + 1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER

# Set column widths
for row in table.rows:
    row.cells[0].width = Cm(1.0)
    row.cells[1].width = Cm(10.5)
    row.cells[2].width = Cm(5.0)

# Header row
header_data = ['№', 'Вопрос', 'Зачем уточнять']
for ci, val in enumerate(header_data):
    cell = table.rows[0].cells[ci]
    cell.text = ''
    p = cell.paragraphs[0]
    run = p.add_run(val)
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), '1A3C6E')
    shading.set(qn('w:val'), 'clear')
    cell._element.get_or_add_tcPr().append(shading)

# Data rows
for ri, (num, question, reason) in enumerate(questions, start=1):
    row = table.rows[ri]
    for ci, val in enumerate([num, question, reason]):
        cell = row.cells[ci]
        cell.text = ''
        p = cell.paragraphs[0]
        run = p.add_run(val)
        run.font.size = Pt(9)
        if ci == 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run.bold = True
        if ri % 2 == 0:
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), 'F2F5FA')
            shading.set(qn('w:val'), 'clear')
            cell._element.get_or_add_tcPr().append(shading)

doc.add_paragraph()

# ---- Closing ----
p = doc.add_paragraph()
run = p.add_run(
    'После получения ответов смогу подготовить:'
    '\n   • учебный план (сетку часов);'
    '\n   • конспект / методическое пособие;'
    '\n   • презентации к каждому модулю;'
    '\n   • раздаточные материалы;'
    '\n   • тестовые задания (итоговые и промежуточные);'
    '\n   • инструкцию для работы в СДО.'
)
run.font.size = Pt(11)

doc.add_paragraph()

p = doc.add_paragraph()
run = p.add_run('Готов ответить на встречные вопросы и предложить предварительную структуру программы.'
                '\nС уважением, [Ваше имя]')
run.font.size = Pt(11)

# ---- Save ----
output_path = r'C:\Users\user\Downloads\for_Cloude Code\prepord-card\Вопросы_заказчику.docx'
doc.save(output_path)
print(f'OK: {output_path}')
print(f'Size: {os.path.getsize(output_path)} B')
